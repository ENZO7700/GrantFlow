"""Export GrantFlow Blueprint Markdown → DOCX + PDF."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from markdown import markdown
from xhtml2pdf import pisa

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "GrantFlow-SK-Blueprint-v2.md"
DOCX_PATH = ROOT / "GrantFlow-SK-Blueprint-v2.docx"
PDF_PATH = ROOT / "GrantFlow-SK-Blueprint-v2.pdf"
HTML_PATH = ROOT / "GrantFlow-SK-Blueprint-v2.html"

CSS = """
@page { size: A4; margin: 1.8cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 10pt; line-height: 1.35; color: #111; }
h1 { font-size: 18pt; color: #0F172A; border-bottom: 1px solid #2563EB; padding-bottom: 4px; }
h2 { font-size: 14pt; color: #0F172A; margin-top: 18px; }
h3 { font-size: 12pt; color: #1E293B; }
h4 { font-size: 11pt; color: #334155; }
code, pre { font-family: Courier; font-size: 8pt; background: #F1F5F9; }
pre { padding: 8px; white-space: pre-wrap; }
table { border-collapse: collapse; width: 100%; margin: 8px 0; font-size: 9pt; }
th, td { border: 1px solid #CBD5E1; padding: 4px 6px; vertical-align: top; }
th { background: #E2E8F0; }
blockquote { border-left: 3px solid #2563EB; margin-left: 0; padding-left: 10px; color: #334155; }
a { color: #2563EB; }
"""


def md_to_html(md_text: str) -> str:
    body = markdown(
        md_text,
        extensions=["tables", "fenced_code", "toc", "sane_lists"],
        output_format="html5",
    )
    return f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"/><style>{CSS}</style></head>
<body>{body}</body></html>"""


def export_pdf(html: str, out: Path) -> None:
    with out.open("wb") as f:
        result = pisa.CreatePDF(html, dest=f, encoding="utf-8")
    if result.err:
        raise RuntimeError(f"PDF export failed with {result.err} errors")


def export_docx(md_text: str, out: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []
    table_buf: list[str] = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = []
        for row in table_buf:
            if re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", row):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            rows.append(cells)
        table_buf = []
        if not rows:
            return
        cols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=cols)
        t.style = "Table Grid"
        for r_idx, row in enumerate(rows):
            for c_idx in range(cols):
                cell = t.cell(r_idx, c_idx)
                cell.text = row[c_idx] if c_idx < len(row) else ""
                if r_idx == 0:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
        doc.add_paragraph("")

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph("\n".join(code_buf))
                for run in p.runs:
                    run.font.name = "Courier New"
                    run.font.size = Pt(8)
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            # peek — if next isn't table, flush
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                flush_table()
            continue

        flush_table()

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line), style="List Number")
        elif line.strip() == "---":
            doc.add_paragraph("—" * 20)
        elif line.strip() == "":
            pass
        else:
            # strip simple markdown bold/links for readability
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            text = re.sub(r"`([^`]+)`", r"\1", text)
            text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
            doc.add_paragraph(text)

        i += 1

    flush_table()
    doc.save(out)


def main() -> int:
    if not MD_PATH.exists():
        print(f"Missing {MD_PATH}", file=sys.stderr)
        return 1
    md_text = MD_PATH.read_text(encoding="utf-8")
    html = md_to_html(md_text)
    HTML_PATH.write_text(html, encoding="utf-8")
    export_docx(md_text, DOCX_PATH)
    export_pdf(html, PDF_PATH)
    print(f"OK DOCX: {DOCX_PATH}")
    print(f"OK PDF:  {PDF_PATH}")
    print(f"OK HTML: {HTML_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
